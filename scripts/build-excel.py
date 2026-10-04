#!/usr/bin/env python3
"""Semester 5 tracker: sheets linked by formulas. Dashboard reads live data."""

from collections import Counter

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins

INK, MUTED, ACCENT, SAGE = "1C1915", "6B6458", "2F4A42", "4F6F62"
PAPER, SURFACE, ELEV, LINE = "F3EFE6", "FBF8F1", "FFFFFF", "DDD4C4"
DONE, WIP, TODO, ACCENT_FG, HIGH = "3D6B4F", "8A6A32", "B9B0A2", "F6F3EA", "E8EFEA"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
F = {
    "paper": PatternFill("solid", fgColor=PAPER),
    "surf": PatternFill("solid", fgColor=SURFACE),
    "elev": PatternFill("solid", fgColor=ELEV),
    "accent": PatternFill("solid", fgColor=ACCENT),
    "sage": PatternFill("solid", fgColor=SAGE),
    "high": PatternFill("solid", fgColor=HIGH),
    "done": PatternFill("solid", fgColor=DONE),
    "wip": PatternFill("solid", fgColor=WIP),
    "todo": PatternFill("solid", fgColor="EEEAE2"),
    "lock": PatternFill("solid", fgColor="EFE8DC"),
}
font_title = Font(name="Calibri", size=26, bold=True, color=INK)
font_h2 = Font(name="Calibri", size=16, bold=True, color=INK)
font_h3 = Font(name="Calibri", size=13, bold=True, color=INK)
font_label = Font(name="Calibri", size=9, bold=True, color=SAGE)
font_muted = Font(name="Calibri", size=10, color=MUTED)
font_body = Font(name="Calibri", size=11, color=INK)
font_white = Font(name="Calibri", size=11, bold=True, color=ACCENT_FG)
font_small_w = Font(name="Calibri", size=9, bold=True, color=ACCENT_FG)
font_phase = Font(name="Calibri", size=18, bold=True, color=INK)
font_kpi = Font(name="Calibri", size=22, bold=True, color=ACCENT)
font_code = Font(name="Consolas", size=9, color=ACCENT)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)
top_left = Alignment(horizontal="left", vertical="top", wrap_text=True)

MONTHS = [
    {"id": "Okt", "label": "Oktober 2026", "phase": "Berani memulai",
     "focus": "Public speaking + social openness",
     "intent": "Bangun fondasi tanpa menunggu percaya diri.", "wpm": 60, "n": 11,
     "question": "Apa yang perlu mulai dilakukan tanpa menunggu percaya diri?",
     "next": "Keberanian bicara, tampil, dan membuat konten jadi dasar kenyamanan.",
     "gsa_weeks": (6, 10), "gsa_target": 10},
    {"id": "Nov", "label": "November 2026", "phase": "Mulai terbiasa",
     "focus": "Conversation + GSA consistency",
     "intent": "Ubah kebiasaan yang baru dimulai menjadi ritme.", "wpm": 70, "n": 9,
     "question": "Bagaimana membuat perilaku Oktober menjadi lebih natural?",
     "next": "Percakapan, public speaking, GSA, English, dan teknis dilakukan konsisten.",
     "gsa_weeks": (11, 14), "gsa_target": 8},
    {"id": "Des", "label": "Desember 2026", "phase": "Menjadi terampil",
     "focus": "Public speaking + technical output",
     "intent": "Skill mulai menghasilkan karya yang terlihat.", "wpm": 80, "n": 9,
     "question": "Bagaimana mengubah kebiasaan menjadi kemampuan yang terlihat?",
     "next": "Komunikasi dan teknis mulai menghasilkan output, portfolio, dan kesiapan interview.",
     "gsa_weeks": (15, 18), "gsa_target": 8},
    {"id": "Jan", "label": "Januari 2027", "phase": "Menerapkan",
     "focus": "GSA completion + real-world communication",
     "intent": "Pakai skill di situasi nyata, finish strong.", "wpm": 90, "n": 6,
     "question": "Apakah skill yang dibangun benar-benar bisa dipakai?",
     "next": "Semua kemampuan dipakai di situasi nyata, termasuk finish GSA dan kesiapan karier.",
     "gsa_weeks": (19, 23), "gsa_target": 10},
]

PATHS = [
    ("Communication", "berani → terbiasa → terampil → natural"),
    ("Social & Mindset", "interaksi → percakapan → relasi"),
    ("GSA & Content", "adaptasi → konsistensi → portfolio"),
    ("Career", "eksplorasi → aplikasi → seleksi"),
    ("Technical", "Excel → SQL → data analysis"),
]
RULES = [
    ("GSA", "2 konten/minggu. Ranking dan reward adalah bonus."),
    ("Kuliah", "Prioritas utama. Jangan dikorbankan demi konten."),
    ("Mental", "Progres lebih penting dari kesempurnaan."),
    ("Public speaking", "Berani → terbiasa → terampil. Anxiety tidak harus 0."),
    ("WPM", "+10 WPM per bulan. Akurasi tetap penting (≥95%)."),
    ("Career", "Eksplorasi → aplikasi → seleksi. Bangun CV + portfolio."),
]
ACH = [
    ("Okt", "Communication", "Public speaking — berani bicara", "Berani menyampaikan pendapat dan berbicara 1–3 menit meskipun masih gugup", "Tidak menghindari kesempatan bicara; minimal 1 sesi latihan/minggu", "Utama", None),
    ("Okt", "Social & Mindset", "Banyakin senyum & ekspresi ramah", "Lebih sering menunjukkan ekspresi ramah saat bertemu/berbicara", "Lebih sering tersenyum dan membuka interaksi", "Sedang", None),
    ("Okt", "Social & Mindset", "Lebih terbuka dengan orang lain", "Mulai sengaja membuka percakapan beberapa kali dalam seminggu", "Ada inisiatif menyapa/bertanya tanpa selalu menunggu", "Utama", None),
    ("Okt", "Social & Mindset", "Kurangi self-judgment", "Menyadari saat mulai menghakimi diri dan menggantinya dengan evaluasi", "Menulis apa yang terjadi dan apa yang bisa diperbaiki, tanpa label negatif", "Utama", None),
    ("Okt", "GSA & Content", "Adaptasi GSA", "Memahami ritme GSA dan mengikuti kegiatan yang diperlukan", "Meeting/briefing penting diikuti; alur kerja dipahami", "Utama", None),
    ("Okt", "GSA & Content", "Konsisten membuat konten", "2 konten per minggu tanpa mengorbankan kuliah dan kesehatan", "Jumlah konten mingguan tercapai secara realistis", "Utama", "gsa:Okt"),
    ("Okt", "English", "Fondasi bahasa Inggris", "Menyelesaikan bagian awal buku dan mulai speaking/listening", "Mampu melakukan percakapan sederhana", "Sedang", None),
    ("Okt", "Technical", "Excel & statistik dasar", "Memahami dasar Excel dan statistik untuk analisis data", "Bisa memakai formula dasar, filter/sort, dan statistik dasar", "Utama", None),
    ("Okt", "Technical", "Typing 60 WPM", "60 WPM dengan akurasi baik", "Mencapai 60 WPM secara konsisten", "Sedang", "wpm:0"),
    ("Okt", "Career", "Eksplorasi magang", "Menentukan bidang dan daftar target magang; siapkan CV/dokumen", "Daftar target + CV awal siap", "Utama", None),
    ("Okt", "Academic", "Sistem kuliah", "Catatan lebih rapi dan materi penting tersimpan di Google Drive", "Catatan mingguan dan file materi terdokumentasi", "Utama", None),
    ("Nov", "Communication", "Public speaking — nyaman", "Presentasi 3–5 menit dan mulai improvisasi tanpa teks penuh", "Bisa berbicara dengan struktur sederhana dan tidak cepat blank", "Utama", None),
    ("Nov", "Social & Mindset", "Mempertahankan percakapan", "Mampu melanjutkan percakapan dan lebih aktif bertanya", "Tidak cepat mengakhiri interaksi karena overthinking", "Utama", None),
    ("Nov", "Social & Mindset", "Self-judgment → evaluasi", "Membedakan 'aku melakukan kesalahan' dari 'aku gagal'", "Kesalahan sosial dinilai sebagai data untuk perbaikan", "Utama", None),
    ("Nov", "GSA & Content", "Konsistensi & gaya konten", "2 konten per minggu dan mulai menemukan gaya komunikasi", "Punya bank ide dan beberapa format yang konsisten", "Utama", "gsa:Nov"),
    ("Nov", "English", "Percakapan sehari-hari", "Lebih lancar dalam speaking/listening", "Tidak terlalu sering menerjemahkan setiap kalimat di kepala", "Sedang", None),
    ("Nov", "Technical", "Excel menengah + SQL dasar", "PivotTable, formula penting, visualisasi; SQL SELECT/WHERE/GROUP BY/JOIN", "Bisa mengerjakan kasus data sederhana", "Utama", None),
    ("Nov", "Technical", "Typing 70 WPM", "70 WPM dengan akurasi baik", "Mencapai 70 WPM secara konsisten", "Sedang", "wpm:1"),
    ("Nov", "Career", "Mulai mendaftar magang", "CV final dan beberapa lamaran terkirim", "Ada bukti aplikasi yang dikirim", "Utama", None),
    ("Nov", "Technical", "Digital product", "Menyelesaikan 1 digital product sederhana", "Ada versi yang dapat ditunjukkan/didemokan", "Sedang", None),
    ("Des", "Communication", "Public speaking — terampil", "Presentasi 5–10 menit, runtut, minim ketergantungan pada teks", "Bisa menjelaskan dan merespons pertanyaan lebih stabil", "Utama", None),
    ("Des", "Social & Mindset", "Lebih natural dalam kelompok", "Aktif berdiskusi dan lebih nyaman berbicara dengan orang baru", "Mulai berinisiatif ikut percakapan/kelompok", "Utama", None),
    ("Des", "Social & Mindset", "Mengurangi takut dinilai", "Tidak terlalu lama memikirkan interaksi yang tidak sempurna", "Lebih cepat kembali ke aktivitas setelah interaksi", "Utama", None),
    ("Des", "GSA & Content", "Pengembangan personal branding", "2 konten per minggu dan gaya konten lebih jelas", "Storytelling/editing/penyampaian makin konsisten", "Utama", "gsa:Des"),
    ("Des", "English", "Fluency bertahap", "Mampu berbicara lebih panjang dan menjelaskan opini sederhana", "Lebih nyaman tanpa terjemahan terus-menerus", "Sedang", None),
    ("Des", "Technical", "Excel + data analysis", "Memakai Excel untuk analisis data sederhana dan statistik dasar", "Ada minimal 1 hasil analisis yang rapi", "Utama", None),
    ("Des", "Technical", "Typing 80 WPM", "80 WPM dengan akurasi baik", "Mencapai 80 WPM secara konsisten", "Sedang", "wpm:2"),
    ("Des", "Career", "Persiapan interview", "CV + portfolio siap; latihan pertanyaan interview", "Mampu menjawab pertanyaan umum lebih terstruktur", "Utama", None),
    ("Des", "Academic", "Arah skripsi", "Memiliki 1–3 kandidat topik yang realistis", "Ada alasan dan arah awal untuk tiap kandidat", "Sedang", None),
    ("Jan", "Communication", "Public speaking — natural dalam praktik", "Berbicara/briefing lebih natural dan tetap berfungsi saat gugup", "Tidak menghindari kesempatan berbicara", "Utama", None),
    ("Jan", "Social & Mindset", "Lebih terbuka & self-judgment terkendali", "Lebih nyaman memulai interaksi dan tidak berlebihan menyalahkan diri", "Fokus pada pembelajaran setelah interaksi", "Utama", None),
    ("Jan", "GSA & Content", "Finish strong, bukan burnout", "Selesaikan GSA sampai 31 Jan dan jaga ritme 2 konten/minggu bila realistis", "Karya terkumpul dan pengalaman siap diceritakan", "Utama", "gsa:Jan"),
    ("Jan", "English", "Practical English", "Menggunakan English dalam situasi nyata", "Conversation lebih spontan dan nyaman", "Sedang", None),
    ("Jan", "Technical", "Typing 90 WPM + portfolio", "90 WPM dengan akurasi baik dan 1 hasil kerja untuk portfolio", "90 WPM + ada output yang bisa ditunjukkan", "Sedang", "wpm:3"),
    ("Jan", "Career", "Siap seleksi magang", "CV/portfolio/interview siap dan strategi aplikasi lebih matang", "Siap mengikuti proses seleksi", "Utama", None),
]
GSA_WEEKS = [
    (1, "1–7 Okt"), (2, "8–14 Okt"), (3, "15–21 Okt"), (4, "22–28 Okt"),
    (5, "29 Okt – 4 Nov"), (6, "5–11 Nov"), (7, "12–18 Nov"), (8, "19–25 Nov"),
    (9, "26 Nov – 2 Des"), (10, "3–9 Des"), (11, "10–16 Des"), (12, "17–23 Des"),
    (13, "24–30 Des"), (14, "31 Des – 6 Jan"), (15, "7–13 Jan"),
    (16, "14–20 Jan"), (17, "21–27 Jan"), (18, "28–31 Jan"),
]
AREAS = [
    ("Communication", "Komunikasi"),
    ("Social & Mindset", "Sosial"),
    ("GSA & Content", "GSA"),
    ("English", "English"),
    ("Technical", "Teknis"),
    ("Career", "Karier"),
    ("Academic", "Kuliah"),
]


def paint(ws, r, c, value=None, font=None, fill=None, align=None, border=None, num=None):
    cell = ws.cell(r, c, value)
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = align
    if border:
        cell.border = border
    if num:
        cell.number_format = num
    return cell


def fill_range(ws, r1, c1, r2, c2, fill):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).fill = fill


def box(ws, r1, c1, r2, c2, fill):
    fill_range(ws, r1, c1, r2, c2, fill)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).border = thin


def merge(ws, r1, c1, r2, c2):
    if r1 != r2 or c1 != c2:
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def color_doughnut(chart, colors):
    series = chart.series[0]
    pts = []
    for i, hexcol in enumerate(colors):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = hexcol
        pts.append(pt)
    series.data_points = pts
    series.graphicalProperties.line.noFill = True


def style_sheet(ws, cols, tab=ACCENT):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins = PageMargins(0.4, 0.4, 0.5, 0.4)
    ws.sheet_properties.tabColor = tab
    for i, w in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.zoomScale = 90


def build():
    wb = Workbook()
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True

    # ── 4 Bulan (sumber data achievement) ────────────────────
    tr = wb.create_sheet("4 Bulan", 0)
    style_sheet(tr, [10, 20, 34, 40, 40, 12, 14, 12, 24, 18])
    fill_range(tr, 1, 1, 100, 10, F["paper"])
    paint(tr, 2, 1, "4 BLOK BULAN  ·  sumber data achievement", font_title, F["paper"], left)
    merge(tr, 2, 1, 2, 9)
    tr.row_dimensions[2].height = 32
    paint(
        tr, 3, 1,
        "Isi kolom Progress (0–100) atau biarkan rumus otomatis (WPM/GSA). Status terhitung dari Progress. Ringkasan, pie, dan sheet lain membaca sheet ini.",
        font_muted, F["paper"], left,
    )
    merge(tr, 3, 1, 3, 9)

    headers = ["Bulan", "Jalur", "Pencapaian", "Target akhir bulan", "Indikator", "Prioritas", "Status", "Progress", "Catatan", "Sumber rumus"]
    month_ranges = {}
    row_by_link = {}
    row = 5
    for m in MONTHS:
        items = [a for a in ACH if a[0] == m["id"]]
        box(tr, row, 1, row, 10, F["accent"])
        merge(tr, row, 1, row, 10)
        paint(tr, row, 1, f"{m['label'].upper()}    ·    {m['phase'].upper()}    ·    {m['focus']}", font_white, F["accent"], left)
        tr.row_dimensions[row].height = 26
        row += 1
        for i, h in enumerate(headers, 1):
            paint(tr, row, i, h, font_small_w, F["sage"], center, thin)
        row += 1
        start = row
        for a in items:
            month, area, title, target, ind, prio, link = a
            fill = F["elev"] if prio == "Utama" else F["surf"]
            for i, v in enumerate([month, area, title, target, ind, prio], 1):
                al = center if i in (1, 6) else top_left
                fnt = Font(name="Calibri", size=10, bold=True, color=ACCENT) if i == 6 and prio == "Utama" else font_body
                paint(tr, row, i, v, fnt, fill, al, thin)
            # H progress
            if link and link.startswith("wpm:"):
                idx = int(link.split(":")[1])
                wrow = 6 + idx
                tgt = MONTHS[idx]["wpm"]
                paint(tr, row, 8, f'=IFERROR(MIN(100,N(WPM!D{wrow})/{tgt}*100),0)', font_code, F["lock"], center, thin, "0")
                paint(tr, row, 10, f"WPM!D{wrow} / {tgt}", font_muted, F["lock"], left, thin)
            elif link and link.startswith("gsa:"):
                mid = link.split(":")[1]
                mm = next(x for x in MONTHS if x["id"] == mid)
                a1, a2, tgt = mm["gsa_weeks"][0], mm["gsa_weeks"][1], mm["gsa_target"]
                paint(tr, row, 8, f'=IFERROR(MIN(100,SUM(GSA!D{a1}:D{a2})/{tgt}*100),0)', font_code, F["lock"], center, thin, "0")
                paint(tr, row, 10, f"GSA!D{a1}:D{a2} / {tgt}", font_muted, F["lock"], left, thin)
            else:
                paint(tr, row, 8, 0, font_body, fill, center, thin, "0")
                paint(tr, row, 10, "input manual", font_muted, fill, left, thin)
            # G status from progress
            paint(tr, row, 7, f'=IF(H{row}>=100,"Selesai",IF(H{row}>0,"Jalan","Belum"))', font_code, F["lock"], center, thin)
            paint(tr, row, 9, "", font_body, fill, left, thin)
            if link:
                row_by_link[link] = row
            tr.row_dimensions[row].height = 44
            row += 1
        end = row - 1
        month_ranges[m["id"]] = (start, end)
        box(tr, row, 1, row, 10, F["high"])
        merge(tr, row, 1, row, 6)
        paint(tr, row, 1, f"Ringkasan {m['label']}", font_h3, F["high"], left, thin)
        paint(tr, row, 7, f'=COUNTIF(G{start}:G{end},"Selesai")&" selesai / "&COUNTA(A{start}:A{end})', font_muted, F["high"], center, thin)
        paint(tr, row, 8, f"=IFERROR(AVERAGE(H{start}:H{end})/100,0)", font_kpi, F["high"], center, thin, "0%")
        paint(tr, row, 9, "", font_body, F["high"], center, thin)
        paint(tr, row, 10, f"G{start}:H{end}", font_muted, F["high"], left, thin)
        tr.row_dimensions[row].height = 24
        row += 2

    last = row
    tr.conditional_formatting.add(
        f"H5:H{last}",
        ColorScaleRule(start_type="num", start_value=0, start_color="EEEAE2",
                       mid_type="num", mid_value=50, mid_color="C4B48A",
                       end_type="num", end_value=100, end_color=DONE),
    )
    tr.conditional_formatting.add(f"G5:G{last}", FormulaRule(formula=['G5="Selesai"'], fill=F["done"], font=font_white))
    tr.conditional_formatting.add(f"G5:G{last}", FormulaRule(formula=['G5="Jalan"'], fill=F["wip"], font=font_white))
    tr.conditional_formatting.add(f"G5:G{last}", FormulaRule(formula=['G5="Belum"'], fill=F["todo"]))
    tr.freeze_panes = "A5"
    tr.column_dimensions["J"].hidden = False

    # ── WPM ──────────────────────────────────────────────────
    wpm = wb.create_sheet("WPM")
    style_sheet(wpm, [20, 22, 14, 14, 12, 16, 16, 16])
    fill_range(wpm, 1, 1, 40, 8, F["paper"])
    paint(wpm, 2, 1, "WPM  ·  tertaut ke Typing di sheet 4 Bulan", font_title, F["paper"], left)
    merge(wpm, 2, 1, 2, 8)
    wpm.row_dimensions[2].height = 32
    paint(wpm, 3, 1, "Isi Aktual WPM. Progress baris Typing (60/70/80/90) di 4 Bulan ikut terhitung. Status memakai target + akurasi.", font_muted, F["paper"], left)
    merge(wpm, 3, 1, 3, 8)
    wh = ["Bulan", "Fase", "Target WPM", "Aktual WPM", "Gap", "Akurasi target", "Aktual akurasi", "Status"]
    for i, h in enumerate(wh, 1):
        paint(wpm, 5, i, h, font_small_w, F["accent"], center, thin)
    for i, m in enumerate(MONTHS):
        r = 6 + i
        box(wpm, r, 1, r, 8, F["elev"])
        paint(wpm, r, 1, m["label"], font_body, F["elev"], left, thin)
        paint(wpm, r, 2, m["phase"], font_body, F["elev"], left, thin)
        paint(wpm, r, 3, m["wpm"], font_h3, F["elev"], center, thin)
        paint(wpm, r, 4, 0, font_body, F["elev"], center, thin)
        paint(wpm, r, 5, f"=IF(D{r}=\"\",C{r},C{r}-D{r})", font_code, F["lock"], center, thin)
        paint(wpm, r, 6, 0.95, font_muted, F["elev"], center, thin, "0%")
        paint(wpm, r, 7, "", font_body, F["elev"], center, thin, "0%")
        paint(wpm, r, 8, f'=IF(OR(D{r}="",D{r}=0),"Belum",IF(AND(D{r}>=C{r},OR(G{r}="",G{r}>=F{r})),"Tercapai","Jalan"))', font_code, F["lock"], center, thin)
        wpm.row_dimensions[r].height = 26
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.title = "WPM: target vs aktual (live)"
    chart.add_data(Reference(wpm, min_col=3, min_row=5, max_col=4, max_row=9), titles_from_data=True)
    chart.set_categories(Reference(wpm, min_col=1, min_row=6, max_row=9))
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 100
    chart.width = 18
    chart.height = 10
    chart.style = 10
    if chart.series:
        chart.series[0].graphicalProperties.solidFill = TODO
        if len(chart.series) > 1:
            chart.series[1].graphicalProperties.solidFill = ACCENT
    wpm.add_chart(chart, "A12")

    # ── GSA ──────────────────────────────────────────────────
    gsa = wb.create_sheet("GSA")
    style_sheet(gsa, [12, 22, 12, 12, 10, 14, 28, 26])
    fill_range(gsa, 1, 1, 45, 8, F["paper"])
    paint(gsa, 2, 1, "GSA  ·  tertaut ke doughnut Ringkasan & baris konten 4 Bulan", font_title, F["paper"], left)
    merge(gsa, 2, 1, 2, 8)
    gsa.row_dimensions[2].height = 32
    paint(gsa, 3, 1, "Isi Aktual tiap minggu. Total, pie, dan progress 'konsisten membuat konten' mengikuti rumus SUM.", font_muted, F["paper"], left)
    merge(gsa, 3, 1, 3, 8)
    gh = ["Minggu", "Periode", "Target", "Aktual", "Gap", "Status", "Ide / topik", "Catatan"]
    for i, h in enumerate(gh, 1):
        paint(gsa, 5, i, h, font_small_w, F["accent"], center, thin)
    for i, (wk, per) in enumerate(GSA_WEEKS):
        r = 6 + i
        fill = F["elev"] if i % 2 == 0 else F["surf"]
        paint(gsa, r, 1, wk, font_body, fill, center, thin)
        paint(gsa, r, 2, per, font_body, fill, left, thin)
        paint(gsa, r, 3, 2, font_body, fill, center, thin)
        paint(gsa, r, 4, 0, font_body, fill, center, thin)
        paint(gsa, r, 5, f"=C{r}-D{r}", font_code, F["lock"], center, thin)
        paint(gsa, r, 6, f'=IF(D{r}>=C{r},"Tercapai",IF(D{r}>0,"Jalan","Belum"))', font_code, F["lock"], center, thin)
        paint(gsa, r, 7, "", font_body, fill, left, thin)
        paint(gsa, r, 8, "", font_body, fill, left, thin)
        gsa.row_dimensions[r].height = 22
    gsa.conditional_formatting.add("F6:F23", FormulaRule(formula=['F6="Tercapai"'], fill=F["done"], font=font_white))
    gsa.conditional_formatting.add("F6:F23", FormulaRule(formula=['F6="Jalan"'], fill=F["wip"], font=font_white))
    paint(gsa, 25, 1, "Total aktual (rumus)", font_h3, F["paper"], left)
    paint(gsa, 25, 4, "=SUM(D6:D23)", font_kpi, F["lock"], center, thin)
    paint(gsa, 26, 1, "Target semester", font_muted, F["paper"], left)
    paint(gsa, 26, 4, 36, font_h3, F["paper"], center)
    paint(gsa, 27, 1, "Progress GSA", font_muted, F["paper"], left)
    paint(gsa, 27, 4, "=IFERROR(D25/D26,0)", font_kpi, F["lock"], center, thin, "0%")

    # ── Rumus (mesin tautan antar-sheet) ─────────────────────
    ms = wb.create_sheet("Rumus")
    style_sheet(ms, [22, 18, 18, 18, 18, 55], SAGE)
    fill_range(ms, 1, 1, 50, 6, F["paper"])
    paint(ms, 1, 1, "MESIN RUMUS  ·  semua visual membaca sel di sini", font_title, F["paper"], left)
    merge(ms, 1, 1, 1, 6)
    ms.row_dimensions[1].height = 30
    paint(ms, 2, 1, "Jangan hapus sheet ini. Ubah data hanya di 4 Bulan / WPM / GSA — angka di bawah mengikuti otomatis.", font_muted, F["paper"], left)
    merge(ms, 2, 1, 2, 6)

    paint(ms, 4, 1, "STATUS KESELURUHAN", font_small_w, F["accent"], center)
    merge(ms, 4, 1, 4, 2)
    paint(ms, 5, 1, "Label", font_small_w, F["sage"], center, thin)
    paint(ms, 5, 2, "Jumlah", font_small_w, F["sage"], center, thin)
    paint(ms, 6, 1, "Selesai", font_body, F["elev"], left, thin)
    paint(ms, 6, 2, '=COUNTIF(\'4 Bulan\'!G:G,"Selesai")', font_code, F["lock"], center, thin)
    paint(ms, 7, 1, "Jalan", font_body, F["elev"], left, thin)
    paint(ms, 7, 2, '=COUNTIF(\'4 Bulan\'!G:G,"Jalan")', font_code, F["lock"], center, thin)
    paint(ms, 8, 1, "Belum", font_body, F["elev"], left, thin)
    paint(ms, 8, 2, '=COUNTIF(\'4 Bulan\'!G:G,"Belum")', font_code, F["lock"], center, thin)
    paint(ms, 9, 1, "Kode", font_muted, F["paper"], left)
    paint(ms, 9, 2, "COUNTIF('4 Bulan'!G:G)", font_muted, F["paper"], left)
    merge(ms, 9, 2, 9, 6)

    paint(ms, 11, 1, "JALUR (pie)", font_small_w, F["accent"], center)
    merge(ms, 11, 1, 11, 4)
    paint(ms, 12, 1, "Jalur", font_small_w, F["sage"], center, thin)
    paint(ms, 12, 2, "Rata progress", font_small_w, F["sage"], center, thin)
    paint(ms, 12, 3, "Jumlah item", font_small_w, F["sage"], center, thin)
    paint(ms, 12, 4, "Nilai pie", font_small_w, F["sage"], center, thin)
    paint(ms, 12, 5, "Kode", font_small_w, F["sage"], center, thin)
    for i, (area, short) in enumerate(AREAS):
        r = 13 + i
        paint(ms, r, 1, short, font_body, F["elev"], left, thin)
        paint(ms, r, 2, f"=IFERROR(AVERAGEIF('4 Bulan'!B:B,\"{area}\",'4 Bulan'!H:H),0)", font_code, F["lock"], center, thin, "0.0")
        paint(ms, r, 3, f"=COUNTIF('4 Bulan'!B:B,\"{area}\")", font_code, F["lock"], center, thin)
        paint(ms, r, 4, f"=IF(SUM($B$13:$B$19)=0,C{r},B{r})", font_code, F["lock"], center, thin, "0.0")
        paint(ms, r, 5, "jika semua 0 → jumlah item; jika ada progres → rata progress", font_muted, F["paper"], left)

    paint(ms, 21, 1, "4 BULAN (kartu dashboard)", font_small_w, F["accent"], center)
    merge(ms, 21, 1, 21, 5)
    paint(ms, 22, 1, "Kode", font_small_w, F["sage"], center, thin)
    paint(ms, 22, 2, "Rata %", font_small_w, F["sage"], center, thin)
    paint(ms, 22, 3, "Selesai", font_small_w, F["sage"], center, thin)
    paint(ms, 22, 4, "Total", font_small_w, F["sage"], center, thin)
    paint(ms, 22, 5, "Label kartu", font_small_w, F["sage"], center, thin)
    paint(ms, 22, 6, "Kode", font_small_w, F["sage"], center, thin)
    for i, m in enumerate(MONTHS):
        r = 23 + i
        s, e = month_ranges[m["id"]]
        paint(ms, r, 1, m["id"], font_body, F["elev"], center, thin)
        paint(ms, r, 2, f"=IFERROR(AVERAGE('4 Bulan'!H{s}:H{e})/100,0)", font_code, F["lock"], center, thin, "0%")
        paint(ms, r, 3, f"=COUNTIF('4 Bulan'!G{s}:G{e},\"Selesai\")", font_code, F["lock"], center, thin)
        paint(ms, r, 4, f"=COUNTA('4 Bulan'!A{s}:A{e})", font_code, F["lock"], center, thin)
        paint(ms, r, 5, f'=C{r}&" / "&D{r}&" selesai"', font_code, F["lock"], left, thin)
        paint(ms, r, 6, f"AVERAGE/COUNTIF 4 Bulan!H{s}:H{e}", font_muted, F["paper"], left)

    paint(ms, 28, 1, "GSA", font_small_w, F["accent"], center)
    merge(ms, 28, 1, 28, 2)
    paint(ms, 29, 1, "Label", font_small_w, F["sage"], center, thin)
    paint(ms, 29, 2, "Nilai", font_small_w, F["sage"], center, thin)
    paint(ms, 30, 1, "Sudah", font_body, F["elev"], left, thin)
    paint(ms, 30, 2, "=SUM(GSA!D6:D23)", font_code, F["lock"], center, thin)
    paint(ms, 31, 1, "Sisa target", font_body, F["elev"], left, thin)
    paint(ms, 31, 2, "=MAX(0,36-B30)", font_code, F["lock"], center, thin)
    paint(ms, 32, 1, "Progress", font_body, F["elev"], left, thin)
    paint(ms, 32, 2, "=IFERROR(B30/36,0)", font_code, F["lock"], center, thin, "0%")
    paint(ms, 33, 1, "Kode", font_muted, F["paper"], left)
    paint(ms, 33, 2, "SUM(GSA!D6:D23) dan 36-konten", font_muted, F["paper"], left)

    paint(ms, 35, 1, "WPM (live ke grafik)", font_small_w, F["accent"], center)
    merge(ms, 35, 1, 35, 3)
    paint(ms, 36, 1, "Bulan", font_small_w, F["sage"], center, thin)
    paint(ms, 36, 2, "Target", font_small_w, F["sage"], center, thin)
    paint(ms, 36, 3, "Aktual", font_small_w, F["sage"], center, thin)
    for i, m in enumerate(MONTHS):
        r = 37 + i
        wr = 6 + i
        paint(ms, r, 1, m["short"] if "short" in m else m["id"], font_body, F["elev"], left, thin)
        paint(ms, r, 2, f"=WPM!C{wr}", font_code, F["lock"], center, thin)
        paint(ms, r, 3, f"=N(WPM!D{wr})", font_code, F["lock"], center, thin)

    wb.defined_names.add(DefinedName(name="StatusSelesai", attr_text="Rumus!$B$6"))
    wb.defined_names.add(DefinedName(name="StatusJalan", attr_text="Rumus!$B$7"))
    wb.defined_names.add(DefinedName(name="StatusBelum", attr_text="Rumus!$B$8"))
    wb.defined_names.add(DefinedName(name="GsaSudah", attr_text="Rumus!$B$30"))
    wb.defined_names.add(DefinedName(name="GsaSisa", attr_text="Rumus!$B$31"))

    # ── Ringkasan ────────────────────────────────────────────
    dash = wb.create_sheet("Ringkasan", 0)
    style_sheet(dash, [3.2, 14, 14, 14, 2.2, 14, 14, 14, 2.2, 14, 14, 14])
    fill_range(dash, 1, 1, 70, 12, F["paper"])
    paint(dash, 2, 2, "OKTOBER 2026  –  JANUARI 2027", font_label, F["paper"], left)
    merge(dash, 2, 2, 2, 12)
    paint(dash, 3, 2, "Planning Semester 5", font_title, F["paper"], left)
    merge(dash, 3, 2, 3, 12)
    dash.row_dimensions[3].height = 36
    paint(dash, 4, 2, "Semua angka dan grafik tertaut rumus ke sheet 4 Bulan, WPM, dan GSA (lewat sheet Rumus).", font_muted, F["paper"], left)
    merge(dash, 4, 2, 4, 12)

    starts = [2, 5, 8, 11]
    for i, m in enumerate(MONTHS):
        c0, c1 = starts[i], starts[i] + 2
        rr = 23 + i
        box(dash, 6, c0, 13, c1, F["surf"])
        merge(dash, 6, c0, 6, c1)
        paint(dash, 6, c0, m["label"].upper(), font_small_w, F["accent"], center)
        dash.row_dimensions[6].height = 22
        merge(dash, 7, c0, 8, c1)
        paint(dash, 7, c0, m["phase"], font_phase, F["surf"], center)
        merge(dash, 9, c0, 9, c1)
        paint(dash, 9, c0, m["focus"], font_muted, F["surf"], center)
        merge(dash, 10, c0, 11, c1)
        paint(dash, 10, c0, m["intent"], font_body, F["surf"], center)
        merge(dash, 12, c0, 12, c1)
        paint(dash, 12, c0, f"=Rumus!B{rr}", font_kpi, F["surf"], center, num="0%")
        merge(dash, 13, c0, 13, c1)
        paint(dash, 13, c0, f"=Rumus!E{rr}", font_muted, F["surf"], center)
        dash.row_dimensions[12].height = 30

    paint(dash, 15, 2, "Status keseluruhan", font_h2, F["paper"], left)
    merge(dash, 15, 2, 15, 4)
    paint(dash, 15, 6, "Progres per jalur", font_h2, F["paper"], left)
    merge(dash, 15, 6, 15, 8)
    paint(dash, 15, 10, "GSA vs target 36", font_h2, F["paper"], left)
    merge(dash, 15, 10, 15, 12)

    ch1 = DoughnutChart()
    ch1.holeSize = 58
    ch1.title = "Status (live)"
    ch1.add_data(Reference(ms, min_col=2, min_row=5, max_row=8), titles_from_data=True)
    ch1.set_categories(Reference(ms, min_col=1, min_row=6, max_row=8))
    ch1.width = 9.2
    ch1.height = 7.2
    color_doughnut(ch1, [DONE, WIP, TODO])
    dash.add_chart(ch1, "B16")

    ch2 = DoughnutChart()
    ch2.holeSize = 58
    ch2.title = "Jalur (live)"
    ch2.add_data(Reference(ms, min_col=4, min_row=12, max_row=19), titles_from_data=True)
    ch2.set_categories(Reference(ms, min_col=1, min_row=13, max_row=19))
    ch2.width = 9.2
    ch2.height = 7.2
    color_doughnut(ch2, [ACCENT, SAGE, WIP, DONE, "8A7A62", MUTED, INK])
    dash.add_chart(ch2, "F16")

    ch3 = DoughnutChart()
    ch3.holeSize = 58
    ch3.title = "GSA (live)"
    ch3.add_data(Reference(ms, min_col=2, min_row=29, max_row=31), titles_from_data=True)
    ch3.set_categories(Reference(ms, min_col=1, min_row=30, max_row=31))
    ch3.width = 9.2
    ch3.height = 7.2
    color_doughnut(ch3, [DONE, TODO])
    dash.add_chart(ch3, "J16")

    paint(dash, 31, 2, '=Rumus!A6&" "&Rumus!B6', font_white, F["done"], center, thin)
    paint(dash, 31, 3, '=Rumus!A7&" "&Rumus!B7', font_white, F["wip"], center, thin)
    paint(dash, 31, 4, '=Rumus!A8&" "&Rumus!B8', font_body, F["todo"], center, thin)
    paint(dash, 31, 10, '=Rumus!B30&" / 36 konten"', font_white, F["accent"], center, thin)
    merge(dash, 31, 10, 31, 12)

    paint(dash, 33, 2, "Lima jalur utama", font_h2, F["paper"], left)
    merge(dash, 33, 2, 33, 12)
    coords = [(2, 3), (4, 5), (6, 7), (8, 9), (10, 12)]
    for i, (name, steps) in enumerate(PATHS):
        c0, c1 = coords[i]
        box(dash, 34, c0, 37, c1, F["elev"])
        merge(dash, 34, c0, 35, c1)
        paint(dash, 34, c0, name, font_h3, F["elev"], center)
        merge(dash, 36, c0, 37, c1)
        paint(dash, 36, c0, steps, font_muted, F["elev"], center)

    paint(dash, 39, 2, "Roadmap", font_h2, F["paper"], left)
    merge(dash, 39, 2, 39, 12)
    for i, m in enumerate(MONTHS):
        c0 = 2 + i * 3
        c1 = c0 + 2
        box(dash, 40, c0, 47, c1, F["surf"])
        merge(dash, 40, c0, 40, c1)
        paint(dash, 40, c0, m["label"].upper(), font_small_w, F["accent"], center)
        merge(dash, 41, c0, 42, c1)
        paint(dash, 41, c0, m["phase"], font_h3, F["surf"], center)
        merge(dash, 43, c0, 45, c1)
        paint(dash, 43, c0, m["question"], font_body, F["surf"], top_left)
        merge(dash, 46, c0, 47, c1)
        paint(dash, 46, c0, m["next"], font_muted, F["surf"], top_left)

    paint(dash, 49, 2, "Aturan main", font_h2, F["paper"], left)
    merge(dash, 49, 2, 49, 12)
    for i, (t, b) in enumerate(RULES):
        c0 = 2 + (i % 3) * 4
        c1 = min(c0 + 3, 12)
        r0 = 50 if i < 3 else 54
        r1 = r0 + 3
        box(dash, r0, c0, r1, c1, F["elev"])
        merge(dash, r0, c0, r0, c1)
        paint(dash, r0, c0, t, font_h3, F["elev"], left)
        merge(dash, r0 + 1, c0, r1, c1)
        paint(dash, r0 + 1, c0, b, font_muted, F["elev"], top_left)

    dash.freeze_panes = "B6"

    # doughnut on GSA sheet too, live
    gch = DoughnutChart()
    gch.holeSize = 60
    gch.title = "GSA vs target (live)"
    gch.add_data(Reference(ms, min_col=2, min_row=29, max_row=31), titles_from_data=True)
    gch.set_categories(Reference(ms, min_col=1, min_row=30, max_row=31))
    gch.width = 10
    gch.height = 8
    color_doughnut(gch, [DONE, TODO])
    gsa.add_chart(gch, "A29")

    help_s = wb.create_sheet("Cara pakai")
    style_sheet(help_s, [3, 100], SAGE)
    fill_range(help_s, 1, 1, 20, 2, F["paper"])
    paint(help_s, 2, 2, "Cara kerja tautan antar-sheet", font_title, F["paper"], left)
    notes = [
        "1. Isi data di 4 Bulan (Progress 0–100), WPM (Aktual WPM), dan GSA (Aktual konten). Jangan ketik di kolom Status — itu rumus.",
        "2. Status di 4 Bulan = IF(Progress>=100,\"Selesai\",IF(Progress>0,\"Jalan\",\"Belum\")).",
        "3. Baris Typing 60/70/80/90 mengambil WPM!D6:D9: Progress = MIN(100, aktual/target*100).",
        "4. Baris konsistensi konten per bulan = SUM aktual GSA minggu bulan itu / target konten bulan itu.",
        "5. Sheet Rumus menghitung COUNTIF/AVERAGEIF/SUM dari ketiga sheet sumber. Pie dan kartu Ringkasan merujuk sel Rumus.",
        "6. Ganti angka di 4 Bulan / WPM / GSA → tekan Enter → Ringkasan, pie, dan kartu bulan berubah (Excel: enable automatic calculation).",
        "7. Buka di Microsoft Excel. Google Sheets mendukung rumusnya; chart Excel kadang perlu dibuat ulang di Sheets.",
        "8. Sheet Rumus adalah mesin kode. Jangan hapus. Named ranges: StatusSelesai, StatusJalan, StatusBelum, GsaSudah, GsaSisa.",
    ]
    for i, t in enumerate(notes):
        paint(help_s, 4 + i, 2, t, font_body, F["paper"], left)
        help_s.row_dimensions[4 + i].height = 26

    # order
    wb._sheets = [dash, tr, wpm, gsa, ms, help_s]
    wb.properties.creator = "Planning Semester 5"
    wb.properties.title = "Planning Semester 5 — linked formulas"

    name = "Planning_Semester_5_Achievement_Tracker_2026_2027.xlsx"
    paths = [
        f"/workspace/artifacts/{name}",
        f"/workspace/public/{name}",
        "/workspace/artifacts/Planning_Semester_5.xlsx",
        "/workspace/public/Planning_Semester_5.xlsx",
        f"/home/workdir/artifacts/{name}",
        "/home/workdir/artifacts/Planning_Semester_5.xlsx",
    ]
    import os
    os.makedirs("/home/workdir/artifacts", exist_ok=True)
    os.makedirs("/workspace/artifacts", exist_ok=True)
    os.makedirs("/workspace/public", exist_ok=True)
    for p in paths:
        wb.save(p)
    print("ok", paths[0], "ranges", month_ranges)


if __name__ == "__main__":
    build()
